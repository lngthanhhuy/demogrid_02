from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Docs/[Tech Demo] Asset Integration and Damage Prototype - Pet Sleeping Mat Oval.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_fixed_width(table_obj):
    tbl = table_obj._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_border(cell, color="D9E2EF", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def status_fill(value):
    text = str(value).lower()
    if text in ("pass", "done", "có", "đạt"):
        return "EAF6ED", "1B5E20"
    if "partial" in text or "warning" in text or "cần" in text or "chờ" in text:
        return "FFF7D6", "7A5A00"
    if "pending" in text or "blocked" in text or "chưa" in text:
        return "FDECEC", "9B1C1C"
    return None, None


def table(rows, widths, header=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    set_table_fixed_width(t)
    for row_idx, row_data in enumerate(rows):
        cells = t.add_row().cells
        for idx, value in enumerate(row_data):
            is_header = header and row_idx == 0
            set_cell_text(
                cells[idx],
                str(value),
                bold=is_header or (idx == 1 and row_idx > 0 and len(row_data) > 2),
                color="FFFFFF" if is_header else None,
                size=8.5 if len(str(value)) > 95 else 9,
            )
            cells[idx].width = Inches(widths[idx])
            if is_header:
                set_cell_shading(cells[idx], "1F4D78")
                set_cell_border(cells[idx], "1F4D78", "8")
            elif idx == 2 and len(row_data) > 2:
                fill, color = status_fill(value)
                if fill:
                    set_cell_shading(cells[idx], fill)
                    for p in cells[idx].paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor.from_string(color)
            elif row_idx % 2 == 0:
                set_cell_shading(cells[idx], "FAFBFC")
        if is_header:
            for cell in cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif len(cells) > 2:
            for p in cells[2].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    return t


def heading(text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def normal(text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p


def evidence_box(title, guidance):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    set_cell_text(cell, f"CHÈN ẢNH / VIDEO MINH CHỨNG TẠI ĐÂY\n\n{guidance}", bold=True, color="555555", size=10)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_shading(cell, "F7FAFD")
    set_cell_border(cell, "8EA9C8", "12")
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), "2200")
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)
    doc.add_paragraph()


OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for name, size, color in [
    ("Heading 1", 15, "2E74B5"),
    ("Heading 2", 12, "2E74B5"),
    ("Heading 3", 11, "1F4D78"),
]:
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].font.bold = True

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("[Tech Demo] Asset Integration & Damage Prototype")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(11, 37, 69)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = subtitle.add_run("Pet Sleeping Mat Oval")
sr.italic = True
sr.bold = True
sr.font.size = Pt(14)
sr.font.color.rgb = RGBColor(46, 116, 181)

table([
    ["Asset", "[FURNITURE - LOT 2] Pet Sleeping Mat Oval", "Assignee", "Danh"],
    ["Priority", "High", "Estimate", "1–2 ngày"],
    ["Ngày cập nhật checklist", "17/08/2026", "Reviewer", "Huy (chờ xác nhận)"],
    ["Unity Project", "demogrid_02", "Render Pipeline", "URP 17.4.0"],
], [1.35, 2.2, 1.35, 2.2], header=False)

heading("1. Mục đích & mục tiêu", 1)
normal(
    "Thực hành quy trình tích hợp model và damage visual prototype trong lúc chờ asset Meat Plant. "
    "Khi Meat Plant được bàn giao, framework chung có thể tái sử dụng — chỉ thay model, data và visual profile."
)
table([
    ["Mục tiêu", "Trạng thái", "Ghi chú"],
    ["Import và kiểm tra model trong Unity", "Pass", "FBX + BC texture đã có trong project."],
    ["Tạo prefab chuẩn", "Pass", "Prefab có root, model child, collider, effect anchor, damage components."],
    ["Xây dựng damage visual prototype tái sử dụng", "Pass", "Framework trong SenCity.Core.DamageVisual."],
    ["Debug controller để test nhanh", "Partial", "Code hoàn tất; cần Canvas UI trong scene hoặc chạy Play để xác nhận bind."],
    ["Chuẩn bị reuse cho Meat Plant", "Pass", "Logic/visual tách rời; không có Growth/Weather/Harvest."],
], [2.2, 1.0, 4.3])

heading("2. Checklist — Import và kiểm tra model", 1)
table([
    ["Hạng mục", "Kết quả", "Bằng chứng / ghi chú"],
    ["Import FBX vào đúng thư mục", "Pass", "Assets/Pet_Sleeping_Mat_Oval/fn_furn_pet_sleeping_mat_oval_01.fbx"],
    ["Import texture Base Color", "Pass", "Assets/Pet_Sleeping_Mat_Oval/fn_furn_pet_sleeping_mat_oval_01_BC.png"],
    ["Naming convention", "Pass", "fn_furn_pet_sleeping_mat_oval_01 theo prefix furniture."],
    ["Scale (import globalScale = 1)", "Pass", "FBX meta globalScale: 1; cần xác nhận visual trong scene."],
    ["Pivot", "Cần xác nhận", "Chưa có báo cáo pivot chính thức; kiểm tra trong Unity Scene view."],
    ["Orientation", "Cần xác nhận", "Model child rotation (0,0,0); xác nhận hướng đặt trên floor."],
    ["Mesh và normals", "Cần xác nhận", "FBX import normalImportMode mặc định; cần inspect mesh trong Editor."],
    ["UV", "Cần xác nhận", "BC texture map tồn tại; chưa có báo cáo UV stretch/overlap."],
    ["Bounds", "Pass", "BoxCollider auto-fit ~ (1.5, 0.05, 1) trên prefab."],
    ["Material slots", "Pass", "1 slot; override material trên prefab child Model."],
    ["Base Color", "Pass", "mat_pet_sleeping_mat_oval.mat gán _BaseMap từ BC texture."],
    ["Normal / ORM map", "N/A", "Asset hiện chỉ có BC; không có Normal/ORM riêng."],
    ["Console Error / Warning", "Cần xác nhận", "Cần mở Unity, Play scene demo và chụp Console sạch lỗi."],
    ["Technical Demo Scene", "Pass", "Assets/_Project/Features/FurniturePlacement/Scenes/DamageVisual_PetSleepingMatDemo.unity"],
], [1.8, 1.0, 4.7])

heading("3. Checklist — Tạo prefab chuẩn", 1)
normal("Prefab path: Assets/_Project/Art/Furniture/PetSleepingMatOval/fn_furn_pet_sleeping_mat_oval_01.prefab")
table([
    ["Yêu cầu", "Kết quả", "Chi tiết"],
    ["Root prefab", "Pass", "Root chứa gameplay/damage components."],
    ["Model child (nested FBX)", "Pass", "Child Model = nested prefab instance; không sửa file FBX import."],
    ["Renderer reference", "Pass", "DamageVisualApplier.targetRenderers + DamageVisualPrefabSetup.targetRenderers."],
    ["Material reference", "Pass", "DamageVisualPrefabSetup.bodyMaterial → mat_pet_sleeping_mat_oval.mat"],
    ["Collider", "Pass", "BoxCollider trên root, auto-fit bounds."],
    ["Damage visual components", "Pass", "DamageState, DamageVisualApplier, DamageStateVisualBinder, DamageVisualPrefabSetup."],
    ["Vị trí gắn effect", "Pass", "Child VFX_Damage tại local Y ≈ 0.07."],
    ["Không sửa trực tiếp FBX", "Pass", "Material override qua prefab instance modification."],
    ["Không có Debug Controller trên prefab production", "Pass", "Debug tách ra scene object Damage Debug Runtime."],
    ["Build tool", "Pass", "Tools/SEN CITY/Damage Visual/Build Pet Sleeping Mat Oval Prefab"],
], [2.0, 1.0, 4.5])

heading("4. Checklist — Damage Visual framework dùng chung", 1)
table([
    ["Yêu cầu", "Kết quả", "Implementation"],
    ["3 state: Normal / Damaged / Destroyed", "Pass", "DamageVisualState enum"],
    ["Data-driven configuration", "Pass", "DamageVisualProfile ScriptableObject"],
    ["Material/texture per state", "Pass", "DamageVisualStateSettings.textureOverride"],
    ["Color/tint fallback", "Pass", "useTintFallback + tintColor (đang dùng cho mat demo)"],
    ["Damage threshold configurable", "Pass", "DamageStateThresholds (25 / 75 default)"],
    ["Renderer targeting", "Pass", "DamageVisualApplier.targetRenderers[]"],
    ["Visual transition", "Pass", "transitionDuration = 0.25s trong profile"],
    ["MaterialPropertyBlock", "Pass", "DamageVisualApplier — reuse 1 MPB, không new Material"],
    ["Texture override", "Pass", "_BaseMap / _MainTex via MPB"],
    ["Damage mask", "Partial", "Hỗ trợ _DamageMask; URP Lit mặc định chưa có property này"],
    ["Không hardcode tên asset", "Pass", "Framework generic trong Core/DamageVisual"],
    ["EditMode tests", "Pass", "Assets/_Project/Tests/EditMode/DamageVisual/DamageStateTests.cs"],
], [2.0, 1.0, 4.5])

heading("5. Checklist — Debug Controller", 1)
table([
    ["Yêu cầu", "Kết quả", "Ghi chú"],
    ["Chuyển Normal", "Pass", "DamageStateDebugController.SetNormal()"],
    ["Chuyển Damaged", "Pass", "DamageStateDebugController.SetDamaged()"],
    ["Chuyển Destroyed", "Pass", "DamageStateDebugController.SetDestroyed()"],
    ["Apply Damage", "Pass", "ApplyDamage() — default +10"],
    ["Reset", "Pass", "ResetDamage()"],
    ["Hiển thị Current State", "Pass", "RefreshUI() cập nhật text"],
    ["Hiển thị Current Damage", "Pass", "Subscribe DamageChanged event"],
    ["Hiển thị Threshold", "Pass", "Hiển thị Damaged/Destroyed threshold"],
    ["Auto-bind UI (không kéo reference)", "Pass", "DamageStateDebugUiResolver tìm theo tên GameObject/label"],
    ["Canvas UI trong demo scene", "Cần xác nhận", "Chưa có Canvas Damage Debug trong scene YAML; cần tạo UI hoặc xác nhận Play test"],
    ["Không sửa Inspector liên tục khi test", "Pass", "Mục tiêu đạt khi UI scene sẵn sàng"],
], [2.0, 1.0, 4.5])

heading("6. Tách phần dùng chung vs riêng (Meat Plant readiness)", 1)
table([
    ["Phân loại", "Thành phần", "Vị trí / asset"],
    ["Dùng chung", "Damage receiver / state logic", "DamageState.cs"],
    ["Dùng chung", "Damage threshold", "DamageStateThresholds.cs"],
    ["Dùng chung", "Visual state controller", "DamageVisualApplier.cs + DamageStateVisualBinder.cs"],
    ["Dùng chung", "Renderer/material binding", "DamageVisualApplier + DamageVisualPrefabSetup"],
    ["Dùng chung", "Debug panel logic", "DamageStateDebugController + DamageStateDebugUiResolver"],
    ["Dùng chung", "Reset flow", "DamageState.Reset()"],
    ["Riêng Furniture (Mat)", "Normal/Damaged/Destroyed visuals", "DamageVisualProfile_PetSleepingMatOval.asset"],
    ["Riêng Furniture (Mat)", "Damage texture thảm", "Chưa có texture riêng; đang dùng tint fallback"],
    ["Riêng Furniture (Mat)", "Furniture configuration", "Prefab + mat_pet_sleeping_mat_oval.mat"],
    ["Dành riêng Meat Plant (chưa làm)", "Healthy/Damaged/Dead", "Ngoài phạm vi task"],
    ["Dành riêng Meat Plant (chưa làm)", "Health Cause: Sun/Wind/Water/Pest", "Ngoài phạm vi task"],
    ["Dành riêng Meat Plant (chưa làm)", "Growth Stage", "Không đưa vào prefab thảm"],
    ["Dành riêng Meat Plant (chưa làm)", "Weather Mapping", "Không implement"],
    ["Dành riêng Meat Plant (chưa làm)", "Harvest Behavior", "Không implement"],
], [1.5, 2.3, 3.7])

heading("7. Definition of Done", 1)
table([
    ["Tiêu chí DoD", "Trạng thái", "Ghi chú"],
    ["Model hiển thị đúng trong Unity", "Cần xác nhận", "Cần screenshot scene demo"],
    ["Scale, pivot, orientation đã kiểm tra", "Partial", "Scale/bounds OK; pivot/orientation cần sign-off visual"],
    ["Material và texture không bị mất", "Pass", "Material asset + BC texture wired"],
    ["Prefab không sửa trực tiếp FBX", "Pass", "Nested prefab + material override"],
    ["Chuyển được 3 trạng thái visual", "Pass", "Logic + applier hoàn chỉnh; cần video demo"],
    ["Reset hoạt động", "Pass", "DamageState.Reset() + unit test"],
    ["Không có Console Error", "Cần xác nhận", "Cần Play test và chụp Console"],
    ["Không tạo material runtime dư thừa", "Pass", "Chỉ dùng sharedMaterial + MPB"],
    ["Damage code không hardcode Pet Sleeping Mat", "Pass", "Generic namespace SenCity.Core.DamageVisual"],
    ["Framework nhận model khác qua config", "Pass", "Profile + prefab setup component"],
    ["Huy review & xác nhận reuse", "Chờ", "Chờ review sau khi bổ sung screenshot/video"],
], [2.4, 1.0, 4.1])

heading("8. Đầu ra (Deliverables)", 1)
table([
    ["Deliverable", "Trạng thái", "Đường dẫn / ghi chú"],
    ["Pet Sleeping Mat Oval đã import", "Done", "Assets/Pet_Sleeping_Mat_Oval/"],
    ["Prefab cấu trúc chuẩn", "Done", "Assets/_Project/Art/Furniture/PetSleepingMatOval/fn_furn_pet_sleeping_mat_oval_01.prefab"],
    ["Technical Demo Scene", "Done", "Assets/_Project/Features/FurniturePlacement/Scenes/DamageVisual_PetSleepingMatDemo.unity"],
    ["Damage Visual Controller dùng chung", "Done", "Assets/_Project/Core/DamageVisual/"],
    ["Damage configuration cho thảm", "Done", "Assets/_Project/Features/FurniturePlacement/Data/DamageVisualProfile_PetSleepingMatOval.asset"],
    ["Debug Controller", "Done", "DamageStateDebugController.cs + UiResolver"],
    ["Asset Validation Report", "Partial", "Checklist này; cần bổ sung ảnh pivot/scale/console"],
    ["Screenshot/video 3 trạng thái", "Chờ", "Chèn vào Phụ lục bên dưới"],
    ["Danh sách phần reuse cho Meat Plant", "Done", "Mục 6 trong tài liệu này"],
    ["Checklist Word", "Done", "File hiện tại"],
], [2.3, 1.0, 4.2])

heading("9. Ngoài phạm vi — xác nhận không implement", 1)
for item in [
    "Sửa mesh hoặc texture của Designer",
    "Growth Stage cho Meat Plant",
    "Weather và Pest",
    "Harvest System",
    "Dùng Furniture state thay Health State chính thức của Meat Plant",
    "Production optimization cuối cùng",
]:
    bullet(f"✓ {item} — không nằm trong scope task này")

heading("10. Kiến trúc & file chính", 1)
table([
    ["File / Asset", "Vai trò"],
    ["DamageVisualState.cs", "Enum Normal / Damaged / Destroyed"],
    ["DamageState.cs", "Logic damage, threshold, reset"],
    ["DamageStateThresholds.cs", "Ngưỡng configurable"],
    ["DamageVisualProfile.cs", "ScriptableObject data-driven"],
    ["DamageVisualStateSettings.cs", "Visual per state"],
    ["DamageVisualApplier.cs", "Apply visual qua MaterialPropertyBlock"],
    ["DamageStateVisualBinder.cs", "Nối logic ↔ visual"],
    ["DamageVisualPrefabSetup.cs", "References: model, material, renderer, collider, VFX anchor"],
    ["DamageStateDebugController.cs", "UI debug bind tự động"],
    ["DamageStateDebugUiResolver.cs", "Tìm Button/Text theo tên"],
    ["DamageVisualPetSleepingMatBuilder.cs", "Editor tool build material/prefab/scene"],
    ["mat_pet_sleeping_mat_oval.mat", "URP Lit + BC texture"],
    ["DamageVisualProfile_PetSleepingMatOval.asset", "Tint damaged/destroyed cho mat demo"],
], [2.5, 4.9], header=True)

heading("11. Hướng dẫn test nhanh", 1)
bullet("Menu: Tools → SEN CITY → Damage Visual → Build Pet Sleeping Mat Oval Demo")
bullet("Mở scene: DamageVisual_PetSleepingMatDemo.unity")
bullet("Play → dùng Damage Debug UI (Normal / Damaged / Destroyed / Apply Damage / Reset)")
bullet("Xác nhận visual đổi màu theo tint trong profile")
bullet("Chạy Test Runner: DamageStateTests (EditMode)")

heading("12. Issue / việc còn lại trước khi Huy review", 1)
table([
    ["Issue", "Mức độ", "Hành động đề xuất"],
    ["Chưa có screenshot/video 3 state", "Cần bổ sung", "Play scene, chụp Normal/Damaged/Destroyed + Reset"],
    ["Console clean chưa có minh chứng", "Cần bổ sung", "Chụp Unity Console khi Play"],
    ["Pivot/orientation chưa sign-off", "Cần xác nhận", "Review visual trong demo scene với art lead"],
    ["Chưa có damage texture riêng", "Chấp nhận", "Prototype dùng tint fallback; texture riêng khi art bàn giao"],
    ["Damage mask chưa dùng được", "Chấp nhận", "Cần custom shader nếu muốn mask"],
    ["Canvas UI debug trong scene", "Cần xác nhận", "Tạo Canvas theo naming convention hoặc xác nhận auto-bind"],
    ["Source FBX vẫn ở drop folder tạm", "Lưu ý", "Chuyển sang _Project/Art khi art được approve theo pipeline"],
    ["Huy review reuse capability", "Chờ", "Gửi tài liệu + demo sau khi bổ sung ảnh"],
], [2.0, 1.0, 4.5])

heading("13. Phụ lục — Khung minh chứng cần chèn", 1)
doc.add_page_break()
evidence_box("A. Import — FBX & Texture trong Project", "Chụp Project window: Pet_Sleeping_Mat_Oval folder với FBX + BC png.")
evidence_box("B. Prefab Hierarchy", "Chụp Hierarchy prefab: Root → Model (nested FBX) → VFX_Damage; Inspector có BoxCollider + Damage components.")
evidence_box("C. Material & Base Color", "Chụp mat_pet_sleeping_mat_oval.mat với _BaseMap wired.")
evidence_box("D. Visual State — Normal", "Play mode, state Normal, screenshot mat gốc.")
evidence_box("E. Visual State — Damaged", "Play mode sau khi bấm Damaged hoặc Apply Damage đủ threshold.")
evidence_box("F. Visual State — Destroyed", "Play mode state Destroyed.")
evidence_box("G. Reset flow", "Screenshot sau Reset về Normal, damage = 0.")
evidence_box("H. Debug UI", "Screenshot Damage Debug panel với state/damage/threshold hiển thị.")
evidence_box("I. Console sạch lỗi", "Screenshot Unity Console khi Play — không Error.")
evidence_box("J. Test Runner", "Screenshot DamageStateTests Pass trong EditMode.")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("[Tech Demo] Pet Sleeping Mat Oval — Damage Prototype Checklist")

doc.save(OUT)
print(OUT.resolve())
