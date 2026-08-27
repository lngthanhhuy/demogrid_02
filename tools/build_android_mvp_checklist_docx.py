from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("SEN_CITY_Android_MVP_Audit_Checklist_Completed.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_fixed_width(table_obj):
    tbl = table_obj._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_border(cell, color="D9E2EF", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def status_fill(value):
    text = str(value).lower()
    if "pass" in text or "done" in text:
        return "EAF6ED", "1B5E20"
    if "warning" in text or "partial" in text or "cần" in text:
        return "FFF7D6", "7A5A00"
    if "blocked" in text or "fail" in text:
        return "FDECEC", "9B1C1C"
    return None, None


def table(rows, widths, header=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    set_table_fixed_width(t)
    for row_idx, row_data in enumerate(rows):
        cells = t.add_row().cells
        for idx, value in enumerate(row_data):
            is_header = header and row_idx == 0
            set_cell_text(
                cells[idx],
                str(value),
                bold=is_header or (idx == 1 and row_idx > 0),
                color="FFFFFF" if is_header else None,
                size=8.5 if len(str(value)) > 90 else 9,
            )
            cells[idx].width = Inches(widths[idx])
            if is_header:
                set_cell_shading(cells[idx], "1F4D78")
                set_cell_border(cells[idx], "1F4D78", "8")
            elif idx == 1:
                fill, color = status_fill(value)
                if fill:
                    set_cell_shading(cells[idx], fill)
                    for p in cells[idx].paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor.from_string(color)
            elif row_idx % 2 == 0:
                set_cell_shading(cells[idx], "FAFBFC")
        if is_header:
            for cell in cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif len(cells) > 1:
            for p in cells[1].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    return t


def heading(text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def normal(text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def evidence_box(title, guidance):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    set_cell_text(cell, f"CHÈN ẢNH MINH CHỨNG TẠI ĐÂY\n\n{guidance}", bold=True, color="555555", size=10)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_shading(cell, "F7FAFD")
    set_cell_border(cell, "8EA9C8", "12")
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), "2300")
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for name, size, color in [
    ("Heading 1", 15, "2E74B5"),
    ("Heading 2", 12, "2E74B5"),
    ("Heading 3", 11, "1F4D78"),
]:
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].font.bold = True

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SEN CITY Unity Project")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(11, 37, 69)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = subtitle.add_run("Checklist tiếp nhận source code, audit project và build Android MVP")
sr.italic = True
sr.font.size = Pt(12)
sr.font.color.rgb = RGBColor(85, 85, 85)

table([
    ["Người thực hiện", "Trần Tiến Danh", "Ngày cập nhật", "13/08/2026"],
    ["Thời gian task", "2 ngày", "Loại task", "Implementation / Build / Verification"],
], [1.45, 2.05, 1.45, 2.05], header=False)

heading("1. Tóm tắt kết quả", 1)
table([
    ["Kết luận nhanh"],
    ["Project Android MVP đã build được APK development sau khi xử lý lỗi duplicate Kotlin dependency. Phần kiểm thử trên thiết bị thật và Logcat đang chờ authorize USB debugging để hoàn tất minh chứng."]
], [7.1])

table([
    ["Hạng mục", "Kết quả", "Ghi chú"],
    ["Repository/branch/commit", "Pass", "Đúng branch codex/sencity-android-mvp, HEAD 8041341."],
    ["Git LFS", "Warning", "Có asset LFS; git lfs fsck báo URP.png không phải pointer."],
    ["Unity project", "Pass có lưu ý", "Unity 6000.4.6f1; scene SenCityMvp có trong build settings."],
    ["Android environment", "Pass", "Unity built-in SDK/NDK/OpenJDK/Gradle sử dụng được để build."],
    ["Android configuration", "Pass có lưu ý", "Cấu hình build ARM64/IL2CPP; package name vẫn là template mặc định."],
    ["Package/plugin compatibility", "Pass sau fix", "Đã xử lý lỗi duplicate Kotlin stdlib trong Gradle."],
    ["Development APK", "Pass", "APK build thành công: Builds/Android/Development/SenCityMvp-Development.apk."],
    ["Install/device smoke test", "Blocked", "ADB thấy thiết bị OZLF7HVG4XJV699H nhưng trạng thái unauthorized."],
    ["Android Logcat", "Blocked", "Chưa thu được logcat vì thiết bị chưa authorize USB debugging."],
], [1.65, 1.15, 4.7])

heading("2. Tiếp nhận repository", 1)
table([
    ["Thông tin", "Giá trị"],
    ["Repository", "https://github.com/lngthanhhuy/demogrid_02.git"],
    ["Workspace", r"C:\Users\h\Documents\SenCity\demogrid_02"],
    ["Branch tiếp nhận", "codex/sencity-android-mvp"],
    ["Commit yêu cầu", "8041341"],
    ["Commit hiện tại", "8041341"],
    ["Scene chính", "Assets/_Project/Production/SenCityMvp.unity"],
    ["Unity Editor", "6000.4.6f1 (0b051c2e5d54)"],
], [2.1, 5.1])

table([
    ["Checklist", "Trạng thái", "Bằng chứng / ghi chú"],
    ["Clone/Pull repository", "Pass", "Remote origin trỏ đến repo GitHub được bàn giao."],
    ["Checkout branch", "Pass", "git branch --show-current = codex/sencity-android-mvp."],
    ["Xác nhận commit", "Pass", "git rev-parse --short HEAD = 8041341."],
    ["Git status", "Warning", "Workspace có thay đổi chưa commit tại thời điểm audit, gồm ProjectSettings, Packages, scene, Assets/Plugins, Resources và Build Profiles."],
    ["Git LFS", "Warning", "git lfs ls-files trả về asset LFS; git lfs fsck báo Assets/TutorialInfo/Icons/URP.png should have been a pointer but was not."],
    ["Merge conflict", "Pass", "Không thấy file conflict marker trong kiểm tra hiện tại."],
], [1.75, 1.0, 4.75])

heading("3. Audit Unity project", 1)
table([
    ["Hạng mục", "Trạng thái", "Ghi chú"],
    ["Project mở bằng đúng Unity version", "Pass", "ProjectVersion.txt ghi Unity 6000.4.6f1."],
    ["Compile/Gradle build", "Pass", "Gradle assembleDebug đã chạy thành công sau fix Kotlin dependency."],
    ["Scene chính", "Pass", "EditorBuildSettings.asset có Assets/_Project/Production/SenCityMvp.unity enabled."],
    ["Missing script trong scene chính", "Pass", "Không thấy m_Script fileID 0 trong SenCityMvp.unity khi quét YAML."],
    ["Missing reference nghiêm trọng", "Warning", "DefaultVolumeProfile.asset có một số m_Script fileID 0; cần xác nhận trong Unity Inspector nếu ảnh hưởng render/post-processing."],
    ["Package Manager resolve", "Pass", "Packages/manifest.json và packages-lock.json có đầy đủ dependency; Gradle đã resolve Android dependency thành công."],
    ["Build Profiles / Scene List", "Pass", "Scene SenCityMvp đang enabled trong EditorBuildSettings."],
], [2.0, 1.0, 4.5])

heading("4. Android build environment", 1)
table([
    ["Thành phần", "Trạng thái", "Phiên bản / đường dẫn", "Ghi chú"],
    ["Android Build Support", "Pass", "Unity AndroidPlayer installed", "Gradle project được tạo trong Library/Bee/Android/Prj/IL2CPP/Gradle."],
    ["Android SDK", "Pass", r"C:\Program Files\Unity\Hub\Editor\6000.4.6f1\Editor\Data\PlaybackEngines\AndroidPlayer\SDK", "Có platforms android-34/35/36, build-tools 36.0.0."],
    ["Android NDK", "Pass", "27.2.12479018", "ndkPath trỏ tới Unity AndroidPlayer/NDK."],
    ["OpenJDK", "Pass", "Unity bundled OpenJDK", "JAVA_HOME trỏ tới Unity AndroidPlayer/OpenJDK."],
    ["Gradle", "Pass", "Gradle 9.1.0 bundled with Unity", "Dùng gradle-launcher-9.1.0.jar."],
    ["ADB", "Partial", r"D:\android\platform-tools\adb.exe", "adb devices thấy thiết bị nhưng unauthorized."],
    ["Android Logcat", "Blocked", "Chưa thu được", "Cần authorize device trước khi lấy logcat."],
], [1.4, 0.9, 2.55, 2.65])

normal("Lưu ý: Gradle vẫn in cảnh báo SDK read-only khi marshalling package.xml, nhưng assembleDebug đã build thành công nên đây là warning môi trường, không phải blocker hiện tại.")

heading("5. Android configuration audit", 1)
table([
    ["Setting", "Value hiện tại", "Nhận xét / rủi ro"],
    ["Company Name", "DefaultCompany", "Nên đổi sang tên công ty/team trước khi phát hành nội bộ/chính thức."],
    ["Product Name", "demogrid_02", "Đang khớp tên project, cần xác nhận tên hiển thị SEN CITY nếu muốn dùng branding MVP."],
    ["Package Name", "com.UnityTechnologies.com.unity.template.urpblank", "Rủi ro cao vì vẫn là package ID template mặc định; nên đổi trước khi test phân phối rộng."],
    ["Version", "0.1.0", "Phù hợp MVP development build."],
    ["Bundle Version Code", "1", "Phù hợp build đầu tiên."],
    ["Minimum API Level", "25", "Android 7.1 trở lên."],
    ["Target API Level", "Automatic / targetSdk 36 trong Gradle output", "Đang dùng SDK mới nhất Unity tìm thấy."],
    ["Scripting Backend", "IL2CPP", "Phù hợp Android ARM64."],
    ["Target Architecture", "ARM64 (AndroidTargetArchitectures: 2)", "Đúng yêu cầu ARM64."],
    ["Build output", "APK", "SenCityMvp-Development.apk được xuất ngoài source code chính trong Builds/Android/Development."],
    ["Keystore", "Chưa cấu hình riêng", "Debug signing được dùng cho development build."],
    ["Custom Manifest", "Không bật", "useCustomMainManifest = 0; useCustomLauncherManifest = 0."],
    ["Custom Gradle Template", "Đã bật baseProjectTemplate", "Bật để ép Kotlin stdlib về 1.8.22 và xử lý duplicate classes."],
    ["Internet Access / Permission", "Chưa xác minh sâu", "Cần test device/logcat nếu MVP có networking."],
], [1.8, 2.1, 3.6])

heading("6. Package và plugin compatibility", 1)
table([
    ["Package/Plugin", "Version", "Android Support", "Issue", "Action"],
    ["com.unity.ai.navigation", "2.0.12", "Có", "Không thấy lỗi build", "Giữ nguyên."],
    ["com.unity.inputsystem", "1.19.0", "Có", "Cần smoke test touch input trên device", "Test sau khi authorize device."],
    ["com.unity.purchasing", "5.4.2", "Có", "Kéo Android Billing 9.0.0, gây xung đột Kotlin với stdlib jdk7/jdk8 1.6.21", "Đã fix bằng Gradle resolution strategy Kotlin 1.8.22."],
    ["com.unity.render-pipelines.universal", "17.4.0", "Có", "DefaultVolumeProfile có m_Script fileID 0 cần xác minh Inspector", "Kiểm tra visual trên device."],
    ["com.unity.test-framework", "1.6.0", "Editor/Test", "Không ảnh hưởng APK runtime", "Giữ nguyên."],
    ["com.unity.timeline", "1.8.12", "Có", "Không thấy lỗi build", "Giữ nguyên."],
    ["com.unity.ugui", "2.0.0", "Có", "Cần smoke test UI trên device", "Test sau khi cài APK."],
    ["com.unity.visualscripting", "1.9.11", "Có nếu dùng graph", "Không thấy lỗi build", "Giữ nguyên."],
    ["Editor packages Rider/VS/Collab", "Nhiều version", "Editor only", "Không ảnh hưởng APK runtime", "Giữ nguyên."],
    ["Assets/Plugins/Android", "Custom Gradle only", "Có", "Không có .aar/.jar/.so native plugin trong Assets khi quét", "Không cần kiểm tra ABI native plugin ở thời điểm này."],
], [1.5, 0.8, 1.15, 2.1, 1.95])

heading("7. Build failure handling và thay đổi đã thực hiện", 1)
table([
    ["Lỗi", "Root cause", "Cách xử lý", "Kết quả"],
    ["Gradle :launcher:checkDebugDuplicateClasses failed", "Duplicate Kotlin classes giữa kotlin-stdlib 1.8.22 và kotlin-stdlib-jdk7/jdk8 1.6.21.", "Thêm Assets/Plugins/Android/baseProjectTemplate.gradle và bật useCustomBaseGradleTemplate để ép kotlin-stdlib* về 1.8.22.", "checkDebugDuplicateClasses Pass; assembleDebug Pass."],
    ["SDK read-only warnings", "Unity SDK nằm trong Program Files nên SDK Manager không ghi package.xml được.", "Ghi nhận warning; không cần fix vì không làm build fail.", "Không còn blocker build."],
], [1.75, 2.0, 2.5, 1.25])

table([
    ["Setting", "Old Value", "New Value", "Reason", "Impact"],
    ["useCustomBaseGradleTemplate", "0", "1", "Cho phép Unity dùng baseProjectTemplate.gradle tùy chỉnh.", "Build Android áp dụng dependency resolution strategy."],
    ["Gradle Kotlin stdlib resolution", "Không ép version", "Ép org.jetbrains.kotlin:kotlin-stdlib* = 1.8.22", "Sửa duplicate classes do dependency kéo nhiều version Kotlin.", "Gradle duplicate class check và assembleDebug thành công."],
], [1.35, 1.1, 1.5, 2.1, 1.45])

heading("8. Android development build", 1)
table([
    ["Hạng mục", "Kết quả", "Bằng chứng / ghi chú"],
    ["Build target", "Pass", "Android, IL2CPP, ARM64."],
    ["Scene included", "Pass", "Assets/_Project/Production/SenCityMvp.unity enabled."],
    ["Gradle task", "Pass", "assembleDebug BUILD SUCCESSFUL."],
    ["APK output", "Pass", r"Builds\Android\Development\SenCityMvp-Development.apk"],
    ["APK size", "Pass", "63,695,325 bytes."],
    ["APK commit policy", "Pass", "Builds/ không nên commit vào Git theo yêu cầu task."],
], [1.6, 1.0, 4.9])

heading("9. Cài APK lên thiết bị Android", 1)
table([
    ["Checklist", "Trạng thái", "Ghi chú"],
    ["Kết nối thiết bị", "Partial", "ADB thấy thiết bị OZLF7HVG4XJV699H."],
    ["Authorize USB debugging", "Blocked", "Thiết bị đang unauthorized; cần bấm Allow USB debugging trên điện thoại."],
    ["Cài APK", "Blocked", "Chưa thể adb install khi device unauthorized."],
    ["Launch application", "Blocked", "Chưa thể launch để xác nhận app mở thành công."],
    ["Device info", "Chờ bổ sung", "Cần ghi Device, Android version, CPU architecture sau khi authorize."],
], [1.75, 1.0, 4.75])

heading("10. Android smoke test", 1)
table([
    ["Test case", "Result", "Ghi chú"],
    ["APK cài thành công", "Blocked", "Chờ authorize thiết bị và adb install."],
    ["Game mở thành công", "Blocked", "Chờ launch trên device."],
    ["Không crash startup", "Blocked", "Cần logcat startup."],
    ["Scene SEN CITY MVP load được", "Blocked", "Cần ảnh/video device."],
    ["UI hiển thị bình thường", "Blocked", "Cần ảnh device."],
    ["Touch Input hoạt động", "Blocked", "Cần thao tác trên device."],
    ["Camera hoạt động", "Blocked", "Cần thao tác trên device."],
    ["Gameplay MVP hoạt động", "Blocked", "Cần thao tác smoke test."],
    ["Không Missing Asset / Missing Script runtime", "Blocked", "Cần logcat và quan sát device."],
    ["Pause/Resume", "Blocked", "Cần test background/resume."],
    ["Đóng và mở lại app", "Blocked", "Cần test relaunch."],
], [2.3, 1.0, 4.2])

heading("11. Android Logcat", 1)
table([
    ["Log cần thu", "Trạng thái", "Ghi chú"],
    ["Startup", "Blocked", "Chưa thu được vì ADB unauthorized."],
    ["Scene load", "Blocked", "Chưa thu được."],
    ["Gameplay", "Blocked", "Chưa thu được."],
    ["Background/Resume", "Blocked", "Chưa thu được."],
    ["Application close", "Blocked", "Chưa thu được."],
], [2.2, 1.0, 4.3])

heading("12. Deliverables", 1)
table([
    ["Deliverable", "Trạng thái", "Đường dẫn / ghi chú"],
    ["Project Handover Verification", "Done", "Có trong checklist này."],
    ["Android Project Audit Report", "Done", "Có trong checklist này."],
    ["Android Environment Verification", "Done", "Có trong checklist này."],
    ["Android Configuration Audit", "Done", "Có trong checklist này."],
    ["Package/Plugin Compatibility Report", "Done", "Có trong checklist này."],
    ["Development APK", "Done", r"Builds\Android\Development\SenCityMvp-Development.apk"],
    ["Build Log", "Partial", "Gradle assembleDebug success; ảnh minh chứng cần chèn bên dưới."],
    ["Android Logcat", "Blocked", "Chờ authorize device."],
    ["Android Smoke Test Report", "Blocked", "Chờ device test."],
    ["Issue/Blocker list", "Done", "Có ở mục 13."],
    ["Đề xuất bước tiếp theo", "Done", "Có ở mục 14."],
], [2.4, 1.0, 4.1])

heading("13. Issue / blocker hiện tại", 1)
table([
    ["Blocker / Issue", "Affected Step", "Evidence", "What was tried", "Required Support"],
    ["ADB device unauthorized", "Install APK, launch app, smoke test, logcat", "adb devices: OZLF7HVG4XJV699H unauthorized", "Đã kiểm tra adb devices", "Mở điện thoại, chấp nhận Allow USB debugging, sau đó chạy lại adb install/logcat."],
    ["Package name template mặc định", "Android configuration / install identity", "com.UnityTechnologies.com.unity.template.urpblank", "Chỉ ghi nhận, chưa đổi vì cần xác nhận naming", "Xác nhận package ID chính thức, ví dụ com.sencity.mvp hoặc theo team."],
    ["Git LFS fsck warning URP.png", "Source handover quality", "URP.png should have been a pointer but was not", "Đã chạy git lfs fsck", "Xác nhận file này có cần quản lý bằng LFS không; sửa pointer nếu cần."],
    ["Workspace có uncommitted changes", "Handover/build reproducibility", "git status --short có modified/untracked files", "Đã ghi nhận", "Review/stage/commit các thay đổi hợp lệ hoặc tách branch con theo quy trình."],
], [1.55, 1.35, 1.55, 1.45, 1.6])

heading("14. Đề xuất bước tiếp theo", 1)
bullet("Authorize thiết bị Android, chạy adb install với APK đã build, launch app và thu logcat.")
bullet("Chèn ảnh minh chứng vào các khung ở phụ lục: build success, APK output, install success, launch screen, smoke test, logcat.")
bullet("Xác nhận và đổi package name khỏi giá trị template mặc định trước khi phát hành nội bộ.")
bullet("Xử lý Git LFS warning URP.png và quyết định các thay đổi đang nằm trong git status.")
bullet("Sau khi smoke test Pass, chuẩn bị Release Build hoặc AAB cho Google Play Internal Testing nếu cần.")

heading("15. Phụ lục ảnh minh chứng cần bổ sung", 1)
doc.add_page_break()
heading("15. Phụ lục ảnh minh chứng cần bổ sung", 1)
normal("Các khung dưới đây được chừa sẵn để bạn dán ảnh minh chứng Android build/device test trực tiếp trong Word.")
evidence_box("A. Android Build Settings / Build Profile", "Dán ảnh Build Profiles hoặc Build Settings có Android, Development Build, ARM64 và scene SenCityMvp.")
evidence_box("B. Gradle build success", "Dán ảnh Console/terminal thể hiện assembleDebug BUILD SUCCESSFUL.")
evidence_box("C. APK output", r"Dán ảnh thư mục Builds\Android\Development có SenCityMvp-Development.apk.")
evidence_box("D. ADB device authorized", "Dán ảnh adb devices sau khi trạng thái chuyển từ unauthorized sang device.")
evidence_box("E. APK install success", "Dán ảnh adb install thành công hoặc màn hình app đã cài trên điện thoại.")
evidence_box("F. App launch / startup", "Dán ảnh màn hình app SEN CITY mở thành công trên thiết bị Android thật.")
evidence_box("G. Smoke test UI / gameplay", "Dán ảnh UI, touch input, camera/gameplay MVP hoạt động.")
evidence_box("H. Android Logcat", "Dán ảnh/log không có crash/error nghiêm trọng trong startup, scene load, gameplay, pause/resume.")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("SEN CITY Android MVP Audit Checklist - Completed")

doc.save(OUT)
print(OUT.resolve())
